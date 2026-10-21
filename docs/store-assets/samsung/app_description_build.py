# Builds the Samsung "App Description" document (template v1.42 structure) from the screenshots and element rects.
# usage: python samsungdoc.py <template.pptx> <shots dir> <out.pptx>
import json, os, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TPL, SHOTS, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
prs = Presentation(TPL)
lst = prs.slides._sldIdLst
for sid in list(lst):
    prs.part.drop_rel(sid.rId); lst.remove(sid)
LAYOUT = next(l for l in prs.slide_layouts if l.name == 'BLANK')
W, H = prs.slide_width, prs.slide_height  # 9144000 x 6858000 (4:3)
GREEN, DARK, GREY = RGBColor(0x39, 0xD3, 0x53), RGBColor(0x20, 0x20, 0x20), RGBColor(0x60, 0x60, 0x60)
M = 400000  # side margin

def slide():
    s = prs.slides.add_slide(LAYOUT)
    for sh in list(s.shapes):
        if sh.is_placeholder: sh._element.getparent().remove(sh._element)
    return s

def box(s, x, y, w, h, text, size=12, bold=False, color=DARK, align=PP_ALIGN.LEFT, fill=None, line=None, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    if fill: tb.fill.solid(); tb.fill.fore_color.rgb = fill
    else: tb.fill.background()
    if line: tb.line.color.rgb = line; tb.line.width = Pt(0.75)
    else: tb.line.fill.background()
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(72000); tf.margin_top = tf.margin_bottom = Emu(36000)
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb

def header(s, section, subtitle=None):
    box(s, M, 200000, W - 2 * M, 450000, section, size=22, bold=True)
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(M), Emu(690000), Emu(W - 2 * M), Emu(25000))
    ln.fill.solid(); ln.fill.fore_color.rgb = GREEN; ln.line.fill.background()
    if subtitle: box(s, M, 760000, W - 2 * M, 320000, subtitle, size=13, bold=True, color=GREEN)

def table(s, x, y, w, rows, col_w, size=10, header=True):
    nrows, ncols = len(rows), len(rows[0])
    t = s.shapes.add_table(nrows, ncols, Emu(x), Emu(y), Emu(w), Emu(280000 * nrows)).table
    for j, cw in enumerate(col_w): t.columns[j].width = Emu(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.text = ''
            tf = c.text_frame; tf.word_wrap = True
            tf.margin_left = tf.margin_right = Emu(50000); tf.margin_top = tf.margin_bottom = Emu(25000)
            r = tf.paragraphs[0].add_run(); r.text = str(val); r.font.size = Pt(size); r.font.bold = (i == 0 and header)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if (i == 0 and header) else DARK
            c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0x40, 0x40, 0x40) if (i == 0 and header) else (RGBColor(0xF4, 0xF4, 0xF4) if i % 2 else RGBColor(0xFF, 0xFF, 0xFF))
    return t

PIC_X, PIC_Y, PIC_W = M, 1150000, 5500000
PIC_H = PIC_W * 1080 // 1920
TBL_X = PIC_X + PIC_W + 150000
TBL_W = W - M - TBL_X

def page(section, subtitle, shot, items, detail):
    s = slide(); header(s, section, subtitle)
    png = os.path.join(SHOTS, shot + '.png'); js = os.path.join(SHOTS, shot + '.json')
    rects = json.load(open(js)) if os.path.exists(js) else {}
    s.shapes.add_picture(png, Emu(PIC_X), Emu(PIC_Y), Emu(PIC_W), Emu(PIC_H))
    rows = [['No', 'Description']]
    for n, (key, desc) in enumerate(items, 1):
        rows.append([str(n), desc])
        r = rects.get(key) if key else None
        if r:
            x, y, w, h = r; sx = PIC_W / 1920
            cx = PIC_X + int((x + w) * sx) - 140000; cy = PIC_Y + int(y * sx) - 70000
            cx = max(PIC_X, min(cx, PIC_X + PIC_W - 240000)); cy = max(PIC_Y, min(cy, PIC_Y + PIC_H - 240000))
            frame = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(PIC_X + int(x * sx)), Emu(PIC_Y + int(y * sx)), Emu(int(w * sx)), Emu(int(h * sx)))
            frame.fill.background(); frame.line.color.rgb = RGBColor(0xFF, 0x3B, 0x3B); frame.line.width = Pt(1.5)
            c = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(cx), Emu(cy), Emu(240000), Emu(240000))
            c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xFF, 0x3B, 0x3B); c.line.fill.background()
            tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; rr = p.add_run(); rr.text = str(n); rr.font.size = Pt(10); rr.font.bold = True; rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    table(s, TBL_X, PIC_Y, TBL_W, rows, [360000, TBL_W - 360000], size=9)
    y = PIC_Y + PIC_H + 150000
    box(s, M, y, W - 2 * M, H - y - 250000, ['Details'] + ['- ' + d for d in detail], size=10, fill=RGBColor(0xF2, 0xF2, 0xF2))
    return s

# 1. Title
s = slide()
box(s, 0, 1900000, W, 900000, 'Easy Xtream Football', size=40, bold=True, align=PP_ALIGN.CENTER)
box(s, 0, 2800000, W, 500000, 'App Description', size=22, color=GREY, align=PP_ALIGN.CENTER)
box(s, 0, 3500000, W, 400000, 'IPTV player for live sport on Samsung Smart TV (Tizen web app)', size=14, color=GREY, align=PP_ALIGN.CENTER)
table(s, 2300000, 4600000, 4544000, [['Seller', 'Nezor (Jorge Martínez)'], ['Contact', 'contact@nezor.es'], ['App ID', 'EXFOOTBALL.EasyXtreamFootball'], ['App version', '1.0.0'], ['Date', '2026-10-22']], [1500000, 3044000], size=11, header=False)

# 2. Revision history
s = slide(); header(s, 'Revision History')
table(s, M, 1100000, W - 2 * M, [['Version', 'Date', 'Description', 'Author'], ['1.0', '2026-10-22', 'Initial document for app version 1.0.0', 'Jorge Martínez']], [1200000, 1400000, 4144000, 1600000], size=11)

# 3. Contents
s = slide(); header(s, 'Contents')
box(s, M + 600000, 1300000, 6000000, 3500000, ['1. UI Structure', '2. Usage Scenario', '3. Menu & function description', '4. Key policy', '5. How to change languages', '6. Notes for QA'], size=18)

# 4. UI structure
s = slide(); header(s, '1. UI Structure', 'Depth navigation (RETURN always goes one level up)')
box(s, M, 1150000, W - 2 * M, 4600000, [
    'Profiles (entry screen)',
    '    ├─ OK on a profile  →  Channel list',
    '    │       ├─ OK on a channel  →  Player  →  ENTER: Player menu (Quality / Audio / Subtitles / Guide / Info / Coffee)',
    '    │       ├─ OK on a folder card  →  Folder (all channels of one brand)  →  Player',
    '    │       └─ magnifier button  →  Filters bar (quality chips, search field, reload)',
    '    ├─ "+ Add profile"  →  Add profile form (Xtream / M3U list / Direct link)  →  Save  →  Channel list',
    '    ├─ gear button  →  Settings (language, clear cache, coffee reminder, open-source licences)',
    '    └─ globe button  →  Language list',
    '',
    'RETURN: Player menu → Player → Channel list → Profiles → exits the app.',
    'Long ENTER on a profile card: Edit / Delete.  Long ENTER on a channel card or in the player: toggle favourite.',
    'Radio stations (audio-only streams) open the Radio view instead of the full-screen player.',
], size=11)

# 5. Usage scenario
s = slide(); header(s, '2. Usage Scenario', 'No login, no sign-up, no payment: the user connects their own IPTV source')
box(s, M, 1150000, W - 2 * M, 1300000, [
    'The app ships with NO channels. On first launch the entry screen is empty except "+ Add profile". The user adds a profile with the data of an IPTV source they already have (an Xtream Codes account, an M3U playlist URL or a direct stream link). Nothing is sent to the developer; the data stays on the TV.',
    'For testing, please use the openly licensed playlists below (no account needed):',
], size=11)
table(s, M, 2500000, W - 2 * M, [
    ['Step', 'Action', 'Expected result'],
    ['1', 'Launch the app', 'Entry screen "Profiles" with the "+ Add profile" card focused'],
    ['2', 'ENTER on "+ Add profile", choose the tab "M3U list"', 'Form with Name and Playlist URL fields'],
    ['3', 'Name: any; Playlist URL: https://cdn.jsdelivr.net/gh/nezor11/easy-xtream-football@main/docs/playlists/review-test.m3u', 'The TV keyboard opens on ENTER; RETURN closes it'],
    ['4', 'ENTER on Save', 'The playlist is downloaded and the Channel list opens (demo videos: Big Buck Bunny, BipBop HLS, Tears of Steel)'],
    ['5', 'ENTER on a channel', 'Full-screen playback; info bar and key legend for a few seconds'],
    ['6', 'ENTER on the picture', 'Player menu: Quality / Audio / Subtitles / Guide / Info / Coffee'],
    ['7', 'RETURN, RETURN, RETURN', 'Back to the Channel list, then Profiles, then the app exits'],
    ['8', 'Optional: add a second M3U profile with https://cdn.jsdelivr.net/gh/nezor11/easy-xtream-football@main/docs/playlists/sports-radio.m3u', 'Public sports radio stations; playing one shows the Radio view (logo, name, "audio only")'],
], [600000, 4600000, 3144000], size=9)

# 6. Menu & function description
page('3. Menu & function description', '3.1 Profiles (entry screen)', '02-profiles', [
    ('title', 'App title and tagline'),
    ('settings', 'Settings button (gear): language, cache, coffee reminder, licences'),
    ('language', 'Language button (globe): pick the app language'),
    ('profile', 'Profile card: ENTER opens its channel list; long ENTER offers Edit / Delete'),
    ('add', '"+ Add profile": opens the form to connect a source'),
    ('hint', 'Hint about long ENTER'),
], ['Entry screen. Focus starts on the first profile (or on "+ Add profile" when there is none). Arrows move the focus, ENTER opens, RETURN exits the app.'])

page('3. Menu & function description', '3.2 Channel list', '05-radios', [
    ('title', 'Profile name and channel count'),
    ('filters', 'Filters button (magnifier): quality filter, search field, reload'),
    ('row', 'Section rows: Live now, Favourites, Recent, Football, More sports (radio stations carry a radio badge)'),
    ('card', 'Channel card: ENTER plays it; long ENTER adds/removes favourite'),
    ('folder', 'Folder card (several channels of one brand): ENTER opens the folder'),
], ['Rows scroll with left/right (they wrap around at the ends); up/down change row. Each card shows the logo or the channel initials, the country and a radio or geo-blocked badge when known.', 'A row with many channels can be opened as a folder; RETURN returns to the list.'])

page('3. Menu & function description', '3.3 Filters bar', '09-filters', [
    ('chip', 'Quality chips: Auto, 4K, FHD, HD, SD (filters the variants shown)'),
    ('search', 'Search field: typing narrows the list live (TV keyboard)'),
    ('reload', 'Reload: downloads the playlist again'),
], ['The magnifier button toggles this bar. RETURN or the magnifier closes it.'])

page('3. Menu & function description', '3.4 Player', '03-player', [
    ('info', 'Channel info bar: name, position in the row, quality, picture size, network speed and seconds buffered'),
    ('legend', 'Key legend (shown for 8 s after opening): ENTER menu, ENTER x2 play/pause, long ENTER favourite, up/down quality, left/right channel'),
], ['Full-screen playback (AVPlay) of the selected channel. Left/right or CH+/- change channel (x2 changes folder), up/down change the quality variant, PLAY/PAUSE or ENTER x2 pause, STOP or RETURN leave the player, INFO hides/shows the info bar.', 'If a stream fails (network, 403/404) a message is shown and the app skips to the next channel; if the source is geo-blocked it says so.'])

page('3. Menu & function description', '3.5 Player menu', '04-player-menu', [
    ('tabs', 'Tabs: Quality, Audio, Subtitles, Guide, Info, Coffee (left/right)'),
    ('option', 'Options of the selected tab: ENTER selects (filled bullet = current)'),
], ['Opened with ENTER over the picture. Quality lists the variants of the channel plus Auto; Audio and Subtitles list the tracks of the stream; Guide shows the programme guide of the channel when the source provides one; Info toggles the info bar; Coffee shows the optional support card (QR code to an external web page; nothing is sold or charged in the app).', 'RETURN closes the menu and returns to the picture.'])

page('3. Menu & function description', '3.6 Radio view', '06-radio-player', [
    ('box', 'Station logo (or a radio glyph)'),
    ('name', 'Station name and "audio only" note'),
    ('info', 'Channel info bar'),
], ['Radio stations (audio-only streams) show this view instead of a black picture. Keys and menu are the same as in the player.'])

page('3. Menu & function description', '3.7 Add profile', '07-add-profile', [
    ('title', 'Screen title'),
    ('tabs', 'Source type tabs: Xtream (server, username, password), M3U list (playlist URL), Direct link (stream URL)'),
    ('field', 'Text fields: ENTER opens the TV keyboard; Done/RETURN on the keyboard returns to the form'),
    ('save', 'Save: checks the data, downloads the channel list and opens it'),
], ['Also used to edit an existing profile (long ENTER on its card). An error message under the form explains a wrong URL or credentials.', 'Test data for QA: see section 2.'])

page('3. Menu & function description', '3.8 Settings', '08-settings', [
    ('title', 'Screen title'),
    ('item', 'Options: Language, Clear channel & guide cache, Buy me a coffee, Show the reminder again, Open-source licences'),
], ['Reached with the gear button on the entry screen. Up/down move, ENTER selects, RETURN returns.'])

# 7. Key policy
s = slide(); header(s, '4. Key policy', 'Samsung standard remote (Smart Remote and full remote)')
table(s, M, 1150000, W - 2 * M, [
    ['Key', 'Lists and forms', 'Player'],
    ['Arrows', 'Move the focus (rows wrap around)', 'Left/right: previous/next channel (x2: folder). Up/down: quality variant'],
    ['ENTER', 'Activate the focused item (long ENTER: favourite / edit)', 'Open the menu (x2: play/pause, long: favourite)'],
    ['RETURN', 'One level up; on the entry screen exits the app', 'Close the menu, else leave the player'],
    ['CH+ / CH-', '-', 'Next / previous channel'],
    ['PLAY / PAUSE / STOP', '-', 'Play-pause toggle / leave the player'],
    ['INFO', '-', 'Show or hide the channel info bar'],
    ['Colour keys, numbers', 'Ignored', 'Ignored'],
    ['HOME / EXIT / VOLUME / MUTE', 'Handled by the TV', 'Handled by the TV'],
    ['Pointer / mouse', 'Not used', 'Not used'],
], [2200000, 3072000, 3072000], size=10)

# 8. Languages
s = slide(); header(s, '5. How to change languages')
box(s, M, 1150000, W - 2 * M, 900000, ['The app follows the TV menu language automatically (24 languages, English fallback). It can also be forced from inside the app:'], size=12)
table(s, M, 2100000, W - 2 * M, [
    ['Step', 'Action'],
    ['1', 'On the entry screen, move the focus to the globe button (top right) and press ENTER'],
    ['2', 'Pick a language in the list and press ENTER; the whole UI changes immediately'],
    ['3', 'Alternatively: gear button → Settings → Language'],
    ['Languages', 'Albanian, Arabic, Basque, Catalan, Chinese (Taiwan), Croatian, Dutch, English, French, Galician, German, Greek, Hindi, Indonesian, Italian, Polish, Portuguese, Romanian, Russian, Serbian, Spanish, Thai, Turkish, Vietnamese'],
], [1200000, 7144000], size=10)

# 9. Notes for QA
s = slide(); header(s, '6. Notes for QA')
box(s, M, 1150000, W - 2 * M, 4800000, [
    '- No login and no account: please use the test playlists of section 2 (openly licensed content served by jsDelivr and public CDNs; no geo-blocking).',
    '- The app is a channel player, not a content service: it has no catalogue of its own and shows only what the user\'s source provides.',
    '- Free app: no in-app purchase, no subscription, no advertising. The "Buy me a coffee" card only shows a QR code to an external web page.',
    '- Live streams cannot be seeked; there is no resume, series or DRM content.',
    '- Quality: "Auto" picks the variant that fits the measured network speed and steps down after stalls; the menu allows manual selection. Variants depend on the source (the test playlists have one variant per item).',
    '- Privilege APIs used: internet, system information (TV model / language), input device key capture (CH+/-, media keys, INFO).',
    '- Tested on a Samsung Smart Monitor M7 LS32AM700U (Tizen 5.5) with a real IPTV subscription (MPEG-TS, H.264, AC-3 / E-AC-3) and with the test playlists.',
    '- Screenshots in this document were taken from the same web app in a desktop browser at 1920x1080 (identical UI).',
    '- Contact for QA questions: contact@nezor.es',
], size=11)

prs.save(OUT); print('saved', OUT, 'slides', len(prs.slides))
